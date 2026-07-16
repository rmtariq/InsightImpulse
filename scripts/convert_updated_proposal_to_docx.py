#!/usr/bin/env python3
"""Convert updated InsightPulse proposal PDF to editable DOCX."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader


PDF = Path("/Users/rmtariq/Downloads/InsightPulse_Proposal_Strategik_2026_V3.1_Singapore_AI_Clinic.pdf")
OUT = Path("/Users/rmtariq/Downloads/InsightPulse_Proposal_Strategik_2026_V3.1_Singapore_AI_Clinic.docx")


def clean_line(line: str) -> str:
    line = re.sub(r"\s+", " ", line or "").strip()
    line = line.replace("—", "-")
    return line


def is_page_marker(line: str) -> bool:
    return bool(re.match(r"^-- \d+ of \d+ --$", line))


def is_header_footer(line: str) -> bool:
    if line.startswith("InsightPulse | Proposal Perkongsian Strategik 2026"):
        return True
    if line.startswith("InsightPulse Sdn Bhd |"):
        return True
    return False


def heading_level(line: str) -> int | None:
    if re.match(r"^\d+\.\s+", line):
        return 1
    if re.match(r"^\d+\.\d+\s+", line):
        return 2
    if line in {
        "PROPOSAL PERKONGSIAN STRATEGIK",
        "InsightPulse AI Platform",
        "Tindakan Seterusnya",
        "HUBUNGI KAMI",
        "Strategic Play untuk Malaysia",
        "Cadangan Program Tambahan: AI Discovery + InsightPulse Pilot",
    }:
        return 1
    return None


def add_para(doc: Document, text: str) -> None:
    lvl = heading_level(text)
    if lvl:
        p = doc.add_heading(text, level=lvl)
        for run in p.runs:
            run.font.color.rgb = RGBColor(0, 51, 102) if lvl == 1 else RGBColor(21, 128, 61)
        return

    if text.startswith("• "):
        p = doc.add_paragraph(text[2:], style="List Bullet")
    else:
        p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05


def convert() -> None:
    reader = PdfReader(str(PDF))
    doc = Document()

    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.7)
    sec.right_margin = Inches(0.7)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(12)

    # Cover page title formatting
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("PROPOSAL PERKONGSIAN STRATEGIK\n")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(0, 51, 102)
    r2 = title.add_run("InsightPulse AI Platform\n")
    r2.bold = True
    r2.font.size = Pt(15)
    r2.font.color.rgb = RGBColor(21, 128, 61)
    r3 = title.add_run("Penyelesaian Kecerdasan Buatan untuk Ekosistem PMKS Malaysia")
    r3.font.size = Pt(11)
    doc.add_paragraph()

    seen_cover = False
    for page_idx, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        lines = [clean_line(x) for x in text.splitlines()]
        lines = [x for x in lines if x and not is_page_marker(x) and not is_header_footer(x)]

        # Skip duplicated cover-title lines because we added formatted cover above.
        if page_idx == 0 and not seen_cover:
            skip = {
                "PROPOSAL PERKONGSIAN STRATEGIK",
                "InsightPulse AI Platform",
                "Penyelesaian Kecerdasan Buatan untuk Ekosistem PMKS Malaysia",
            }
            lines = [x for x in lines if x not in skip]
            seen_cover = True

        for line in lines:
            add_para(doc, line)

        if page_idx != len(reader.pages) - 1:
            doc.add_page_break()

    # Add a small editable note at the end.
    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("DOCX editable version generated from V3.1 PDF.")
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 116, 139)

    doc.save(str(OUT))


if __name__ == "__main__":
    convert()
    print(f"Saved: {OUT}")
