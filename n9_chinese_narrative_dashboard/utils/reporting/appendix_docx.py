"""Analytical appendix — DOCX writer."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from utils.report_formatting import docx_set_col_widths, docx_style_table_header


def write_appendix_docx(payload: dict[str, Any], path: Path) -> Path:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(f"Lampiran Analitik — {payload['state']}", level=0)
    doc.add_paragraph(f"{payload['record_count']:,} rekod | Analisis terperinci untuk penyelidik & fact checker")

    doc.add_heading("Metodologi", level=1)
    for line in payload.get("methodology") or []:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Suara Komuniti (petikan)", level=1)
    for v in payload.get("voice_quotes") or []:
        doc.add_paragraph(f"[{v['platform']}] {v['issue']} — {v['original'][:200]}")
        doc.add_paragraph(f"BM: {v['translation_bm'][:200]} | URL: {v.get('url', '—')}")

    for title, key in [("Taburan Sentimen", "sentiment"), ("Platform", "platform"), ("Isu", "issues"), ("Bahasa", "languages")]:
        rows = payload.get("tables", {}).get(key) or []
        if not rows:
            continue
        doc.add_heading(title, level=1)
        t = doc.add_table(rows=1 + len(rows), cols=2)
        t.style = "Table Grid"
        t.rows[0].cells[0].text = "Kategori"
        t.rows[0].cells[1].text = "Bilangan"
        docx_style_table_header(t.rows[0].cells)
        for i, (a, b) in enumerate(rows):
            t.rows[i + 1].cells[0].text = a
            t.rows[i + 1].cells[1].text = b

    doc.add_heading("Sampel Hantaran (50 teratas)", level=1)
    hdr = ["Platform", "Isu", "Sentimen", "Petikan", "URL"]
    posts = payload.get("all_posts_sample") or []
    if posts:
        pt = doc.add_table(rows=1 + len(posts), cols=5)
        pt.style = "Table Grid"
        docx_set_col_widths(pt, [0.8, 1.2, 0.7, 2.5, 1.8])
        for j, h in enumerate(hdr):
            pt.rows[0].cells[j].text = h
        docx_style_table_header(pt.rows[0].cells)
        for i, row in enumerate(posts):
            for j, val in enumerate(row):
                pt.rows[i + 1].cells[j].text = val
                for run in pt.rows[i + 1].cells[j].paragraphs[0].runs:
                    run.font.size = Pt(7)

    doc.add_heading("Ringkasan Harian (Brief)", level=1)
    for line in (payload.get("brief") or "").splitlines():
        if line.strip().startswith("- "):
            doc.add_paragraph(line.strip()[2:], style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line.strip())

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
