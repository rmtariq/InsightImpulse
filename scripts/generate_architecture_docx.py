#!/usr/bin/env python3
"""Generate InsightPulse architecture DOCX — PRN N9 + Naratif Cina & India."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "InsightPulse_Architecture_PRN_N9_Naratif_Cina_India.docx"

NAVY = RGBColor(0, 51, 102)
GREEN = RGBColor(21, 128, 61)
SLATE = RGBColor(30, 41, 59)
GREY = RGBColor(100, 116, 139)


def set_run(run, *, bold=False, size=10, color=None, mono=False):
    run.bold = bold
    run.font.name = "Consolas" if mono else "Arial"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def h1(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        set_run(r, bold=True, size=15, color=NAVY)


def h2(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=2)
    for r in p.runs:
        set_run(r, bold=True, size=12, color=GREEN)


def h3(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=3)
    for r in p.runs:
        set_run(r, bold=True, size=11, color=SLATE)


def para(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    for r in p.runs:
        set_run(r, size=10, color=SLATE)


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    for r in p.runs:
        set_run(r, size=10, color=SLATE)


def code_block(doc: Document, text: str) -> None:
    for line in text.strip().splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line)
        set_run(r, size=8.5, color=SLATE, mono=True)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, head in enumerate(headers):
        hdr[i].text = head
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        for p in hdr[i].paragraphs:
            for r in p.runs:
                set_run(r, bold=True, size=9, color=NAVY)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for r in p.runs:
                    set_run(r, size=8.5, color=SLATE)
    if widths:
        for row in table.rows:
            for idx, w in enumerate(widths):
                row.cells[idx].width = Inches(w)
    doc.add_paragraph()


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("InsightPulse Architecture Design\n")
    set_run(r, bold=True, size=20, color=NAVY)
    r = p.add_run("PRN N9 Dashboard + Naratif Cina & India Dashboard\n")
    set_run(r, bold=True, size=14, color=GREEN)
    r = p.add_run("A to Z System Architecture")
    set_run(r, size=12, color=SLATE)
    doc.add_paragraph()
    for label, val in [
        ("Platform", "InsightPulse — Malaysian Political Intelligence"),
        ("Scope", "PRN Negeri Sembilan (N9) + Community Narrative (Cina & India)"),
        ("States Covered", "Negeri Sembilan (36 DUN), Johor (56 DUN), Melaka (28 DUN)"),
        ("Tarikh", "Julai 2026"),
        ("Klasifikasi", "Internal"),
    ]:
        p = doc.add_paragraph()
        r = p.add_run(f"{label}: ")
        set_run(r, bold=True, size=10, color=NAVY)
        r = p.add_run(val)
        set_run(r, size=10)
    doc.add_page_break()


def build_doc() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)

    add_title(doc)

    # ── 1. Platform Overview ──
    h1(doc, "1. Platform Overview")
    para(
        doc,
        "InsightPulse is a Malaysian political intelligence platform built around CSV/JSON-first data pipelines. "
        "It powers three operational surfaces for PRN Negeri Sembilan (N9) and a dedicated community-narrative "
        "stack for Chinese & Indian voters across N9, Johor, and Melaka.",
    )

    h2(doc, "1.1 High-Level Architecture Flow")
    code_block(
        doc,
        """
Data Sources (Apify, RSS, SPR Excel, Culaan Forms, Economic Feeds)
        ↓
Processing Pipeline (Crawl → Malay BERT Sentiment → Analyze → Combine → Build)
        ↓
Dashboard Surfaces (Static HTML, War Room LIVE, Streamlit, React, Standalone HTML)
        """,
    )

    h2(doc, "1.2 Tech Stack")
    add_table(
        doc,
        ["Layer", "Technology"],
        [
            ["Primary API", "FastAPI (web_backend/simple_app.py) — port 8001"],
            ["War Room API", "Python ThreadingHTTPServer (scripts/prn_digital_cula_server.py) — port 8080"],
            ["Narrative Dashboard", "Streamlit + Plotly (n9_chinese_narrative_dashboard/) — port 8501"],
            ["PRN Intel UI", "React 19 + TypeScript + Vite + Tailwind + Recharts (frontend/)"],
            ["War Room UI", "Vanilla HTML/CSS/JS + Leaflet maps"],
            ["ML/NLP", "HuggingFace Malay BERT, scikit-learn (SVM/RF/DT), OpenAI/Anthropic"],
            ["Data Store", "CSV-first (primary); optional SQLite via SQLAlchemy"],
        ],
        widths=[1.4, 4.8],
    )

    h2(doc, "1.3 Project Registry")
    para(doc, "All work is scoped under registered projects in data/projects/_registry.json:")
    add_table(
        doc,
        ["Project ID", "Scope"],
        [
            ["PRN_N9", "Negeri Sembilan — 36 DUN seats"],
            ["PRN_Johor", "Johor — 56 DUN seats"],
            ["PRN_Melaka", "Melaka — 28 DUN seats"],
            ["Cina_Narative_2026", "Harakat/Syurga issue — Chinese voter sentiment"],
            ["pas_break_2026", "PAS–Bersatu split monitoring"],
        ],
        widths=[1.8, 4.4],
    )

    doc.add_page_break()

    # ── 2. Data Architecture ──
    h1(doc, "2. Data Architecture (A → Z)")

    h2(doc, "2.1 Top-Level Folder Structure")
    code_block(
        doc,
        """
InsightPulse/
├── backend/              # Crawlers, ML, DB models, report generation
├── web_backend/          # Production FastAPI (10-platform crawl + analyze)
├── frontend/             # React PRN monitoring dashboard
├── n9_chinese_narrative_dashboard/   # Streamlit Cina & India app
├── scripts/              # 100+ pipeline/build/crawl scripts
└── data/
    ├── smart_crawlers/   # Raw crawl per platform
    ├── analyzed/         # Sentiment + emotion per platform
    ├── combined/         # Merged cross-platform CSVs
    └── projects/
        └── political/
            ├── PRN/
            │   ├── PRN_N9/          ← N9 project root
            │   ├── PRN_Johor/
            │   ├── PRN_Melaka/
            │   ├── digital_cula/    ← Live field submissions
            │   └── warroom_live/    ← Live action tracking
            └── Cina_Narative_2026/  ← Harakat/Syurga focused crawl
        """,
    )

    h2(doc, "2.2 Per-Project Data Layout (PRN_N9)")
    code_block(
        doc,
        """
PRN_N9/
├── metadata.json
├── crawls/
│   ├── news/              # Multilingual news (BM/EN/ZH/Tamil)
│   ├── news_analyzed/     # Tier A ML sentiment
│   ├── chinese_news/      # Chinese media (NS-filtered)
│   ├── indian_narrative/  # Tamil/India community crawls
│   ├── social/            # FB/TikTok/X
│   └── adun_seeds/        # ADUN social seeds
├── reference/             # Queries, seeds, JSON summaries
├── reports/
│   ├── dashboards/        # prn-negeri-sembilan-dashboard.html
│   ├── war_room/          # n9_war_room.json, WhatsApp digest
│   │   └── prototype/     # War Room LIVE app
│   └── seats/             # 36 DUN analytics
└── master/                # Merged datasets
        """,
    )

    h2(doc, "2.3 Universal Data Flow")
    code_block(
        doc,
        """
User query / seeds / URLs
        ↓
Crawlers (Apify, RSS, direct scrape)
        ↓
data/smart_crawlers/{platform}/
        ↓
Batch Sentiment Processor (Malay BERT)
        ↓
data/analyzed/Sentiment_Emotion_{PLATFORM}_*.csv
        ↓
Merge → data/combined/Combined_*.csv
        ↓
Copy to project → PRN_N9/crawls/
        ↓
Build scripts → JSON/CSV bundles
        ↓
Dashboards (HTML / Streamlit / War Room)
        """,
    )

    doc.add_page_break()

    # ── 3. PRN N9 Dashboard ──
    h1(doc, "3. PRN N9 Dashboard — Three Surfaces")
    para(doc, "PRN N9 has three dashboard surfaces, each serving a different operational need.")

    h2(doc, "3A. Static Election Dashboard (HTML)")
    bullet(doc, "Output: data/projects/political/PRN/PRN_N9/reports/dashboards/prn-negeri-sembilan-dashboard.html")
    bullet(doc, "Generator: scripts/generate_prn_n9_dashboard.py")
    bullet(doc, "Daily pipeline: scripts/run_prn_n9_daily.sh")

    h3(doc, "Daily Pipeline Steps")
    add_table(
        doc,
        ["Step", "Script", "Output"],
        [
            ["1", "crawl_prn_n9_news_multilingual.py", "PRN_N9/crawls/news/"],
            ["2", "analyze_prn_n9_news.py", "PRN_N9/crawls/news_analyzed/"],
            ["3", "crawl_prn_chinese_news.py", "chinese_news (optional)"],
            ["4", "n9_war_room.py", "war_room JSON + WhatsApp digest"],
            ["5", "generate_prn_n9_dashboard.py", "HTML dashboard"],
        ],
        widths=[0.5, 2.5, 3.2],
    )

    h3(doc, "Data Sources Consumed")
    bullet(doc, "36 DUN seat analytics (reports/seats/n9_seats_export.json)")
    bullet(doc, "SPR Excel (JITP_2026/Master_File/PRN2023_NegeriSembilan_v2.xlsx)")
    bullet(doc, "PAS-Break socmed (pas_break_2026/reports/pas_break_dashboard_data.json)")
    bullet(doc, "Economic indicators (reference/economic_n9.json)")
    bullet(doc, "ML predictions (n9_ml_predictions.json)")
    bullet(doc, "Chinese narrative overlay (Cina_Narative_2026/reports/cina_narrative_analytics.json)")

    h3(doc, "Features")
    bullet(doc, "Seat map (36 DUN), coalition scenarios, sentiment trends")
    bullet(doc, "Economic stress signals, Chinese narrative intel overlay, war room actions")

    h2(doc, "3B. War Room LIVE — Primary Operational Dashboard")
    bullet(doc, "URL: http://localhost:8080/?module=command&state=N9")
    bullet(doc, "Start: bash scripts/start_prn_warroom_live.sh")
    bullet(doc, "UI: data/projects/political/PRN/PRN_N9/reports/war_room/prototype/index.html")
    bullet(doc, "Server: scripts/prn_digital_cula_server.py (port 8080)")

    h3(doc, "9 Modules (Sidebar Navigation)")
    add_table(
        doc,
        ["#", "Module", "Purpose"],
        [
            ["1", "Command Center", "Strategic overview, coalition summary, KPIs"],
            ["2", "DUN & Tindakan", "36-seat map, demographics, drill-down drawer"],
            ["3", "Action Center", "Owner, due date, priority, status tracking"],
            ["4", "Pulse Digital", "Online signal detection"],
            ["5", "Rapid Response", "Alerts + response actions"],
            ["6", "Daily Briefing", "7am auto-briefing from action data"],
            ["7", "Polling Day Mode", "T-72h activation UI"],
            ["8", "Naratif Cina & India", "Community narrative monitoring"],
            ["9", "Culaan Digital", "Field form, PDM review, socmed seeds"],
        ],
        widths=[0.4, 1.6, 4.2],
    )

    h3(doc, "Production Data Bundle")
    para(doc, "Built by scripts/build_warroom_production_bundle.py → synced to prototype/data/:")
    add_table(
        doc,
        ["File", "Content"],
        [
            ["n9_production_bundle.json", "Master production bundle"],
            ["warroom_dun_N9_production.json", "36 DUN with demographics, scenarios"],
            ["n9_war_room.json", "War room playbook"],
            ["socmed_by_dun_N9.json", "Socmed aggregated per DUN"],
            ["n9_ml_predictions.json", "ML seat predictions"],
            ["n9_narrative_community.json", "Chinese + Indian narrative (Module #8)"],
            ["cula_hub_N9.json", "Digital Culaan KPIs"],
        ],
        widths=[2.2, 4.0],
    )

    h3(doc, "Live Persistence (Operator Edits)")
    add_table(
        doc,
        ["Path", "What it stores"],
        [
            ["digital_cula/N9/submissions.csv", "Field Culaan form submissions"],
            ["warroom_live/N9/action_status.json", "Action Center status updates"],
            ["warroom_live/N9/custom_actions.json", "Custom operator actions"],
        ],
        widths=[2.8, 3.4],
    )

    h3(doc, "Key API Endpoints (port 8080)")
    add_table(
        doc,
        ["Endpoint", "Purpose"],
        [
            ["POST /api/cula/submit", "Digital Culaan field submissions"],
            ["GET /api/cula/hub", "Cula KPI aggregation"],
            ["GET/POST /api/action-center/*", "War room action tracking"],
            ["GET /api/pulse-digital", "Digital signal detection"],
            ["GET /api/ml/predict/prn_n9", "ML seat analytics"],
            ["POST /api/content/generate", "Content Studio copy generation"],
        ],
        widths=[2.5, 3.7],
    )

    h3(doc, "Deep Links")
    bullet(doc, "Naratif Cina & India: http://localhost:8080/?module=narrative&state=N9")
    bullet(doc, "Culaan Digital: http://localhost:8080/?module=cula&state=N9")
    bullet(doc, "Action Center: http://localhost:8080/?module=action&state=N9")

    h2(doc, "3C. React PRN Intel Dashboard")
    bullet(doc, "Path: frontend/src/features/prn-monitoring/")
    bullet(doc, "Route: /insightpulse/prn-monitoring/ (served from War Room at port 8080 after build)")
    bullet(doc, "Data modes: Manifest JSON, Excel workbook upload, Mock fallback")
    bullet(doc, "Build: bash scripts/build_prn_react_dashboard.sh")

    doc.add_page_break()

    # ── 4. Naratif Cina & India ──
    h1(doc, "4. Naratif Cina & India Dashboard — Four Surfaces")
    para(
        doc,
        "The community narrative stack monitors Chinese and Indian voter sentiment across "
        "N9 (36 DUN), Johor (56 DUN), and Melaka (28 DUN).",
    )

    h2(doc, "4.1 Narrative Architecture Flow")
    code_block(
        doc,
        """
Crawl Sources (Chinese Media, Tamil/India, Social, Harakat/Syurga)
        ↓
build_dataset.py → community.py classifier → keyword_classifier.py
        ↓
alert_engine.py + action_engine.py
        ↓
4 Dashboard Surfaces (Streamlit, War Room Modul #8, Standalone HTML, Agentic LLM Brief)
        """,
    )

    h2(doc, "4.2 Central Dataset")
    bullet(doc, "Master CSV: n9_chinese_narrative_dashboard/data/prn_chinese_narrative_posts_3negeri.csv")
    bullet(doc, "Builder: n9_chinese_narrative_dashboard/utils/build_dataset.py")

    h3(doc, "Community Classification (utils/community.py)")
    add_table(
        doc,
        ["Community", "Detection"],
        [
            ["chinese", "zh script, Chinese media platforms, query IDs S1/S2/S3"],
            ["indian", "crawl_community=india, Tamil script, Tamil media domains"],
            ["cross", "BM/EN general posts"],
            ["general", "Everything else"],
        ],
        widths=[1.2, 5.0],
    )

    h2(doc, "4A. Streamlit Narrative Dashboard (Detailed Operations)")
    bullet(doc, "Start: cd n9_chinese_narrative_dashboard && .venv/bin/streamlit run app.py")
    bullet(doc, "Port: 8501")

    h3(doc, "Pages")
    add_table(
        doc,
        ["Page", "File", "Focus"],
        [
            ["Ringkasan (home)", "app.py", "KPIs, status HIJAU/MERAH, executive summary"],
            ["Naratif Cina", "pages/1_Naratif_Cina.py", "Chinese community sentiment, issues, top posts"],
            ["Naratif India", "pages/2_Naratif_India.py", "Indian/Tamil community monitoring"],
            ["Peta Isu", "pages/3_Peta_Isu.py", "District/DUN hotspot map"],
            ["Pusat Tindakan", "pages/4_Pusat_Tindakan.py", "Action kanban + response copy"],
            ["Data & Laporan", "pages/5_Data_dan_Laporan.py", "Data table + PDF/DOCX export"],
        ],
        widths=[1.4, 2.0, 3.0],
    )

    h3(doc, "Key Utilities")
    bullet(doc, "utils/keyword_classifier.py — issue clustering (Cost of Living, DAP Performance, etc.)")
    bullet(doc, "utils/alert_engine.py — risk scoring")
    bullet(doc, "utils/action_engine.py — action recommendations")
    bullet(doc, "utils/reporting/ — executive PDF/DOCX, response kits (BM/中文/Tamil)")
    bullet(doc, "utils/translator.py — BM translation for display")

    h2(doc, "4B. War Room Module #8 (Embedded)")
    bullet(doc, "JSON output: prototype/data/n9_narrative_community.json")
    bullet(doc, "Builder: scripts/build_warroom_narrative_json.py")
    bullet(doc, "Daily update: scripts/update_narrative.sh")

    h3(doc, "Update Commands")
    add_table(
        doc,
        ["Command", "Description"],
        [
            ["./scripts/update_narrative.sh pagi", "Chinese news + Indian crawl + publish (~5 min)"],
            ["./scripts/update_narrative.sh petang", "Same for evening slot"],
            ["./scripts/update_narrative.sh", "Publish only (~30 sec, no crawl)"],
            ["./scripts/update_narrative.sh penuh", "Full news + India crawl"],
            ["./scripts/update_narrative.sh laporan", "Generate PDF/DOCX reports"],
        ],
        widths=[2.8, 3.4],
    )

    h3(doc, "Crawl Scripts Invoked")
    bullet(doc, "scripts/crawl_prn_chinese_daily.py — free Chinese news (Google News + RSS, no Apify)")
    bullet(doc, "scripts/crawl_prn_indian_narrative.py — Tamil/India narrative")

    h2(doc, "4C. Standalone HTML Dashboard (Johor + N9)")
    bullet(doc, "Builder: scripts/build_naratif_cina_india_standalone_dashboard.py")
    bullet(doc, "Output: data/projects/political/PRN/reports/naratif_cina_india/naratif_cina_india_johor_n9_dashboard.html")
    bullet(doc, "Side-by-side Chinese vs Indian analytics per state, with theme hits (PH/DAP, PAS/PN, MN/BN, boycott, Harakat, cost of living)")

    h2(doc, "4D. Agentic LLM Analysis Layer")
    bullet(doc, "Script: scripts/analyze_naratif_agentic.py")
    bullet(doc, "RAG corpus: PRN_N9/reference/n9_pn_faction_rag.json + Claude/OpenAI")
    bullet(doc, "Outputs: naratif_agentic_brief.json, naratif_cadangan_tindakan.json, state-specific DOCX action plans")

    h2(doc, "4E. Cina_Narative_2026 Project (Harakat/Syurga Issue)")
    bullet(doc, "Path: data/projects/political/Cina_Narative_2026/")
    bullet(doc, "Purpose: Focused crawl on Harakat/Shawn Loh/syurga tak bau narrative affecting Chinese voter sentiment toward PAS/BN")
    bullet(doc, "Dashboard: data/projects/political/Cina_Narative_2026/reports/cina_narrative_johor_dashboard.html")
    bullet(doc, 'Crawl queries: "罗盛年", "Shawn Loh", "Harakatdaily", "syurga", "天堂", PAS+BN combinations')

    doc.add_page_break()

    # ── 5. Full System Architecture ──
    h1(doc, "5. Full System Architecture")
    code_block(
        doc,
        """
┌─────────────────────────────────────────────────────────────────┐
│                        DATA SOURCES                              │
│  Apify (FB/IG/X/TikTok) │ Google News RSS │ Direct Scrape       │
│  Field Culaan Forms │ Excel SPR Seat Data │ Chinese/Tamil Media │
└──────────────────────────────┬──────────────────────────────────┘
                               ↓
┌─────────────────────────────────────────────────────────────────┐
│                    PROCESSING PIPELINE                             │
│  smart_crawlers/ → Malay BERT Sentiment → analyzed/ → combined/ │
│  → PRN_N9/crawls/ → build_dataset.py / build_warroom_*.py       │
└──────────────────────────────┬──────────────────────────────────┘
                               ↓
┌─────────────────────────────────────────────────────────────────┐
│                        API LAYER                                 │
│  FastAPI :8001 (Crawl + Analyze) │ War Room Server :8080        │
└──────────────────────────────┬──────────────────────────────────┘
                               ↓
┌─────────────────────────────────────────────────────────────────┐
│                    DASHBOARD SURFACES                            │
│  Static HTML │ War Room LIVE (9 Modules) │ Streamlit (6 Pages)  │
│  React PRN Intel │ Standalone HTML Johor+N9                      │
└─────────────────────────────────────────────────────────────────┘
        """,
    )

    # ── 6. Startup & Deployment ──
    h1(doc, "6. Startup & Deployment")
    add_table(
        doc,
        ["Command", "What it starts", "Port"],
        [
            ["./start_full_app.sh", "Full InsightPulse (crawl + analyze)", "8001"],
            ["bash scripts/start_prn_warroom_live.sh", "PRN War Room LIVE + APIs", "8080"],
            ["cd n9_chinese_narrative_dashboard && .venv/bin/streamlit run app.py", "Narrative Streamlit", "8501"],
            ["cd frontend && npm run dev", "React PRN Intel (dev)", "5173"],
        ],
        widths=[3.0, 2.2, 0.8],
    )

    h2(doc, "Daily Automation")
    code_block(
        doc,
        """
# PRN N9 daily pipeline
bash scripts/run_prn_n9_daily.sh

# Narrative Cina & India (2× daily recommended)
./scripts/update_narrative.sh pagi
./scripts/update_narrative.sh petang
        """,
    )

    # ── 7. Key Configuration Files ──
    h1(doc, "7. Key Configuration Files")
    add_table(
        doc,
        ["File", "Purpose"],
        [
            ["data/projects/_registry.json", "All project IDs and metadata"],
            [".env", "API keys (OpenAI, Apify, Anthropic)"],
            ["scripts/prn_paths.py", "State path resolution"],
            ["docs/prd/war-room-prd.md", "War Room V2.0 product requirements"],
            ["n9_chinese_narrative_dashboard/.streamlit/config.toml", "Streamlit theme"],
        ],
        widths=[3.0, 3.2],
    )

    # ── 8. Summary ──
    h1(doc, "8. Summary")
    add_table(
        doc,
        ["Dashboard", "Type", "Primary Use", "Update Frequency"],
        [
            ["PRN N9 Static HTML", "Self-contained HTML", "Executive overview, seat map, coalition scenarios", "Daily (run_prn_n9_daily.sh)"],
            ["War Room LIVE", "SPA + API server :8080", "Operational war room — 9 modules incl. Culaan, Action Center, Narrative", "Real-time + daily bundle refresh"],
            ["Streamlit Naratif", "Streamlit :8501", "Deep Chinese/Indian community analysis, reports, action kits", "2× daily (update_narrative.sh)"],
            ["React PRN Intel", "React SPA", "Excel-based seat intel, scenario modeling", "On Excel upload"],
            ["Standalone HTML", "Self-contained HTML", "Johor+N9 side-by-side Chinese vs Indian", "On build script run"],
        ],
        widths=[1.3, 1.2, 2.2, 1.5],
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("\n— End of Document —")
    set_run(r, size=9, color=GREY)

    return doc


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_doc()
    doc.save(OUT)
    print(f"✅ DOCX generated: {OUT}")


if __name__ == "__main__":
    main()
