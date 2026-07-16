"""Executive 4-page report — DOCX writer."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from utils.report_formatting import docx_set_col_widths, docx_style_table_header


def write_executive_docx(payload: dict[str, Any], path: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    navy = RGBColor(30, 58, 95)

    # ── Cover ──
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(payload.get("org", "InsightPulse"))
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = navy

    h = doc.add_heading(payload["title"], level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(payload["subtitle"])
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in [
        f"Negeri: {payload['state']}",
        f"Tempoh laporan: {payload['data_period']}",
        f"Dijana: {payload['generated_at']}",
        f"Data dikemas kini: {payload['data_updated']}",
        f"Versi: {payload['version']}",
        f"Klasifikasi: {payload['classification']}",
    ]:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # ── Page 1: Ringkasan Eksekutif ──
    doc.add_heading("Ringkasan Eksekutif — Baca Dalam 60 Saat", level=1)
    st = payload["status"]
    p = doc.add_paragraph()
    p.add_run(f"{st['code']} — {st['label']}: ").bold = True
    p.add_run(st["sentence"])

    doc.add_heading("Petunjuk Utama", level=2)
    kpi = doc.add_table(rows=2, cols=4)
    kpi.style = "Table Grid"
    items = list(payload["kpis"].items())
    for i, (lbl, val) in enumerate(items[:8]):
        r, c = divmod(i, 4)
        kpi.rows[r].cells[c].text = f"{lbl}\n{val}"

    if payload["sentiment"].get("warning"):
        doc.add_paragraph(payload["sentiment"]["warning"])

    charts = payload.get("charts") or {}
    for key in ("sentiment", "issues"):
        if charts.get(key):
            doc.add_picture(charts[key], width=Inches(5.8))

    for title, key in [
        ("A. Apa Sedang Berlaku", "happening"),
        ("B. Mengapa Ia Penting", "important"),
        ("C. Keputusan Diperlukan Hari Ini", "decisions"),
    ]:
        doc.add_heading(title, level=2)
        for item in payload["bullets"].get(key, []):
            doc.add_paragraph(item, style="List Bullet")

    if payload["reliability"].get("warning"):
        doc.add_paragraph(payload["reliability"]["warning"])

    doc.add_page_break()

    # ── Page 2: Lima Isu Utama ──
    doc.add_heading("Lima Isu Utama", level=1)
    hdr = ["#", "Isu", "Lokasi", "Hantaran", "% Neg", "Interaksi", "Trend", "Mengapa Penting", "Prioriti"]
    rows = payload.get("top5_issues") or []
    if rows:
        t2 = doc.add_table(rows=1 + len(rows), cols=len(hdr))
        t2.style = "Table Grid"
        docx_set_col_widths(t2, [0.3, 1.2, 0.9, 0.6, 0.5, 0.7, 0.7, 2.0, 0.9])
        for j, h in enumerate(hdr):
            t2.rows[0].cells[j].text = h
        docx_style_table_header(t2.rows[0].cells)
        for i, row in enumerate(rows):
            cells = t2.rows[i + 1].cells
            vals = [str(row["rank"]), row["issue"], row["location"], str(row["count"]),
                    f"{row['neg_pct']}%", f"{row['engagement']:,}", row["trend"], row["why"], row["priority"]]
            for j, v in enumerate(vals):
                cells[j].text = v
                for run in cells[j].paragraphs[0].runs:
                    run.font.size = Pt(8)
    else:
        doc.add_paragraph("Tiada isu strategik melepasi penapis (Belum Diklasifikasikan dikecualikan).")

    doc.add_page_break()

    # ── Page 3: Papan Tindakan ──
    doc.add_heading("Papan Tindakan 24 Jam hingga 7 Hari", level=1)
    for act in payload.get("action_board") or []:
        doc.add_paragraph(f"{act['priority']} — {act['issue_location']}", style="List Bullet")
        for lbl, key in [
            ("Apa Berlaku", "what_happened"), ("Tindakan Khusus", "action"), ("Hasil Diperlukan", "output"),
            ("Pasukan", "team"), ("Tarikh Akhir", "deadline"), ("Status Bukti", "evidence"),
            ("KPI", "kpi"), ("Status", "status"),
        ]:
            doc.add_paragraph(f"{lbl}: {act.get(key, '—')}")
        doc.add_paragraph("")

    doc.add_page_break()

    # ── Page 4: Risiko & Susulan ──
    doc.add_heading("Risiko, Pengesahan dan Susulan", level=1)
    doc.add_heading("A. Risiko Utama", level=2)
    for r in payload.get("top_risks") or []:
        doc.add_paragraph(
            f"• {r['reason']} | {r['location']} | {r['source']} | Bukti: {r['evidence']} | Kelulusan: {r['approval']}",
            style="List Bullet",
        )
    if not payload.get("top_risks"):
        doc.add_paragraph("Tiada risiko Tinggi/Kritikal dalam tempoh data.")

    doc.add_heading("B. Dakwaan Memerlukan Pengesahan", level=2)
    for c in payload.get("claims") or []:
        doc.add_paragraph(f"• {c['summary']} | URL: {c['url']} | Pegawai: {c['officer']} | {c['deadline']}")

    doc.add_heading("C. Pelan Masa", level=2)
    for phase, detail in payload.get("timeline") or []:
        doc.add_paragraph(f"{phase}: {detail}", style="List Bullet")

    doc.add_heading("D. Laporan Seterusnya", level=2)
    nx = payload.get("next_report") or {}
    doc.add_paragraph(f"Tarikh: {nx.get('datetime', '—')}")
    doc.add_paragraph(f"P1 belum selesai: {nx.get('open_p1', 0)} | Pengesahan belum selesai: {nx.get('open_verify', 0)}")
    doc.add_paragraph(f"Pasukan: {nx.get('team', '—')}")

    doc.add_heading("Pengesahan", level=2)
    for field in payload.get("signoff") or []:
        doc.add_paragraph(f"{field}: _________________________")

    doc.add_paragraph("")
    doc.add_paragraph("Lampiran Analitik tersedia sebagai fail berasingan.", style="Intense Quote")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
