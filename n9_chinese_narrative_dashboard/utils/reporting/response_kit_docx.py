"""Kit Respons Naratif — DOCX writer."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from utils.report_formatting import docx_style_table_header


def write_response_kit_docx(payload: dict[str, Any], path: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    navy = RGBColor(30, 58, 95)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(payload.get("org", "InsightPulse"))
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = navy

    h = doc.add_heading(payload["title"], level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(payload["subtitle"])
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tag = doc.add_paragraph(payload.get("tagline", ""))
    tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in [
        f"Negeri: {payload['state']}",
        f"Dijana: {payload['generated_at']}",
        f"Versi: {payload['version']}",
        payload["classification"],
        f"{payload['record_count']:,} rekod analisis",
    ]:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(payload.get("disclaimer", ""), style="Intense Quote")
    doc.add_page_break()

    doc.add_heading("Aliran Kerja", level=1)
    for phase, detail in payload.get("workflow") or []:
        doc.add_paragraph(f"{phase}: {detail}", style="List Bullet")

    for item in payload.get("items") or []:
        doc.add_page_break()
        doc.add_heading(
            f"{item['priority']} — {item['issue']} ({item['location']})",
            level=1,
        )

        doc.add_heading("1. Apa Isu", level=2)
        doc.add_paragraph(f"Naratif: {item.get('narrative_summary', '—')}")
        doc.add_paragraph(f"Tindakan operasi: {item.get('operational_action', '—')}")
        doc.add_paragraph(f"Output: {item.get('output', '—')} | KPI: {item.get('kpi', '—')}")
        src = item.get("source") or {}
        doc.add_paragraph(f"Sumber rujukan: {src.get('platform', '—')} | {src.get('url', '—')}")
        if src.get("text") and src.get("text") != "—":
            doc.add_paragraph(f"Petikan: {src['text'][:250]}")

        doc.add_heading("2. Copy Pos Utama (Draf)", level=2)
        doc.add_paragraph(f"Sumber: {item.get('copy_source', 'template')}")
        doc.add_paragraph("BM — Post ringkas:")
        doc.add_paragraph(item.get("copy_bm_social") or item.get("copy_bm_short", ""))
        doc.add_paragraph("中文 — Post ringkas:")
        doc.add_paragraph(item.get("copy_zh_social") or item.get("copy_zh_short", ""))
        doc.add_paragraph(f"Hashtag: {' '.join(item.get('hashtags') or [])}")
        doc.add_paragraph("BM — Kenyataan penuh:")
        doc.add_paragraph(item.get("copy_bm_long", ""))
        doc.add_paragraph("中文 — Kenyataan penuh:")
        doc.add_paragraph(item.get("copy_zh_long", ""))

        cp = item.get("counter_plan") or {}
        doc.add_heading("3. Naratif Balas — Cara Balas, Platform & Copy", level=2)

        doc.add_paragraph("A. Langkah edar (ikut urutan):", style="List Bullet")
        for i, step in enumerate(cp.get("steps") or [], 1):
            doc.add_paragraph(f"{i}. {step}")

        doc.add_paragraph("B. Copy balas komen (jika ada hantaran negatif):")
        doc.add_paragraph(f"BM: {cp.get('reply_comment_bm', '—')}")
        doc.add_paragraph(f"中文: {cp.get('reply_comment_zh', '—')}")

        doc.add_paragraph("C. Pos mengikut platform (salin & edit [pautan] sebelum hantar):")
        posts = cp.get("platform_posts") or []
        if posts:
            pt = doc.add_table(rows=1, cols=4)
            pt.style = "Table Grid"
            for j, hdr in enumerate(["Platform", "Masa", "% Data", "Copy siap pos"]):
                pt.rows[0].cells[j].text = hdr
            docx_style_table_header(pt.rows[0].cells)
            for pp in posts:
                row = pt.add_row().cells
                row[0].text = pp.get("platform", "")
                row[1].text = pp.get("timing", "")
                row[2].text = f"{pp.get('data_share_pct', 0)}%"
                row[3].text = pp.get("copy", "")
        else:
            for plat in item.get("platforms") or []:
                doc.add_paragraph(
                    f"• {plat.get('platform')} ({plat.get('timing')}): {plat.get('format', '')}",
                    style="List Bullet",
                )

        doc.add_paragraph("D. JANGAN:")
        doc.add_paragraph(cp.get("dont", "—"))
        doc.add_paragraph(f"Nada: {cp.get('tone', '—')}")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
