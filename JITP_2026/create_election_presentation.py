#!/usr/bin/env python3
"""
Create PowerPoint Presentation for:
Big Data & Analytic: Strategi Pilihan Raya dan Taktikal
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Slide 1: Title Slide
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    
    # Title
    title_box = slide1.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1))
    title_frame = title_box.text_frame
    title_frame.text = "Big Data & Analytic"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(54)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(0, 51, 102)
    title_para.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    subtitle_box = slide1.shapes.add_textbox(Inches(0.5), Inches(3.5), Inches(9), Inches(1))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = "Strategi Pilihan Raya dan Taktikal"
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.font.size = Pt(36)
    subtitle_para.font.color.rgb = RGBColor(192, 0, 0)
    subtitle_para.alignment = PP_ALIGN.CENTER
    
    # Author
    author_box = slide1.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(9), Inches(0.5))
    author_frame = author_box.text_frame
    author_frame.text = "InsightPulse Analytics Platform"
    author_para = author_frame.paragraphs[0]
    author_para.font.size = Pt(20)
    author_para.font.italic = True
    author_para.alignment = PP_ALIGN.CENTER
    
    # Slide 2: Agenda
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    title2 = slide2.shapes.title
    title2.text = "Agenda"
    content2 = slide2.placeholders[1]
    tf2 = content2.text_frame
    tf2.text = "Pengenalan Big Data dalam Politik"
    
    p2 = tf2.add_paragraph()
    p2.text = "Kepentingan Analytics dalam Pilihan Raya"
    p2.level = 0
    
    p3 = tf2.add_paragraph()
    p3.text = "Platform InsightPulse - Tools Monitoring"
    p3.level = 0
    
    p4 = tf2.add_paragraph()
    p4.text = "Strategi Taktikal Menggunakan Data"
    p4.level = 0
    
    p5 = tf2.add_paragraph()
    p5.text = "Case Study & Contoh Aplikasi"
    p5.level = 0
    
    p6 = tf2.add_paragraph()
    p6.text = "Kesimpulan & Q&A"
    p6.level = 0
    
    # Slide 3: Pengenalan Big Data dalam Politik
    slide3 = prs.slides.add_slide(prs.slide_layouts[1])
    title3 = slide3.shapes.title
    title3.text = "Big Data dalam Politik: Era Baru Kempen"
    content3 = slide3.placeholders[1]
    tf3 = content3.text_frame
    tf3.text = "✅ Volume Data Besar-besaran"
    
    p = tf3.add_paragraph()
    p.text = "Jutaan posting media sosial setiap hari"
    p.level = 1
    
    p = tf3.add_paragraph()
    p.text = "10 platform: Facebook, Instagram, X, TikTok, YouTube, LinkedIn, News, Lowyat, Shopee, Lazada"
    p.level = 1
    
    p = tf3.add_paragraph()
    p.text = "✅ Kelajuan & Masa Nyata (Real-time)"
    p.level = 0
    
    p = tf3.add_paragraph()
    p.text = "Sentiment berubah dalam masa nyata"
    p.level = 1
    
    p = tf3.add_paragraph()
    p.text = "Crisis detection dalam minit"
    p.level = 1
    
    p = tf3.add_paragraph()
    p.text = "✅ Kepelbagaian Data"
    p.level = 0
    
    p = tf3.add_paragraph()
    p.text = "Text, images, videos, engagement metrics"
    p.level = 1
    
    p = tf3.add_paragraph()
    p.text = "Multilingual: Malay, English, Chinese"
    p.level = 1
    
    # Save
    prs.save('BigData_Election_Strategy_Part1.pptx')
    print("✅ Part 1 created!")
    return prs

if __name__ == "__main__":
    create_presentation()
