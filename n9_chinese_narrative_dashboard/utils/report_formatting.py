"""Shared professional formatting for PDF/DOCX narrative reports."""
from __future__ import annotations

from typing import Any

# Brand palette
NAVY = "#1e3a5f"
SLATE = "#64748b"
LIGHT_BG = "#f1f5f9"
TIER_COLORS = {"P1": "#dc2626", "P2": "#f59e0b", "P3": "#2563eb"}


def _esc(text: str) -> str:
    return (
        str(text or "—")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def priority_actions_from_insights(ins: dict[str, Any]) -> list[dict[str, str]]:
    rows = ins.get("priority_actions") or []
    if rows:
        return rows
    hdr = ins.get("action_table_header") or []
    out: list[dict[str, str]] = []
    for row in ins.get("action_table") or []:
        d = {hdr[i]: str(row[i]) if i < len(row) else "" for i in range(len(hdr))}
        out.append({
            "rank": d.get("P", ""),
            "priority_tier": d.get("Tier", "P3"),
            "issue": d.get("Isu", ""),
            "location": d.get("Lokasi", ""),
            "platform": d.get("Platform", ""),
            "score": d.get("Skor", ""),
            "action": d.get("Tindakan", ""),
            "team": d.get("Pasukan", ""),
            "due": d.get("Tarikh", ""),
            "evidence": d.get("Bukti", ""),
        })
    return out


# ── PDF helpers ──────────────────────────────────────────────────────────────

def pdf_styles():
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet

    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "RptTitle",
            parent=base["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=16,
            textColor=NAVY,
            alignment=TA_CENTER,
            spaceAfter=10,
        ),
        "subtitle": ParagraphStyle(
            "RptSub",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=10,
            textColor=SLATE,
            alignment=TA_CENTER,
            spaceAfter=6,
        ),
        "h1": ParagraphStyle(
            "RptH1",
            parent=base["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=13,
            textColor=NAVY,
            spaceBefore=14,
            spaceAfter=8,
        ),
        "h2": ParagraphStyle(
            "RptH2",
            parent=base["Heading3"],
            fontName="Helvetica-Bold",
            fontSize=11,
            textColor=NAVY,
            spaceBefore=8,
            spaceAfter=4,
        ),
        "body": ParagraphStyle(
            "RptBody",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=10,
            leading=14,
            alignment=TA_JUSTIFY,
        ),
        "bullet": ParagraphStyle(
            "RptBullet",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=10,
            leading=14,
            leftIndent=12,
            spaceBefore=2,
            spaceAfter=2,
        ),
        "snap": ParagraphStyle(
            "RptSnap",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=10,
            leading=15,
        ),
        "cell": ParagraphStyle(
            "RptCell",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=8,
            leading=10,
            alignment=TA_LEFT,
        ),
        "cell_hdr": ParagraphStyle(
            "RptCellHdr",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=10,
            textColor="#ffffff",
            alignment=TA_LEFT,
        ),
        "cell_bold": ParagraphStyle(
            "RptCellBold",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=10,
            alignment=TA_LEFT,
        ),
    }


def pdf_wrapped_table(rows: list[list[str]], col_widths: list[float], styles: dict) -> Any:
    from reportlab.platypus import Paragraph, Table, TableStyle
    from reportlab.lib import colors

    data = []
    for i, row in enumerate(rows):
        data.append([
            Paragraph(_esc(c), styles["cell_hdr"] if i == 0 else styles["cell"])
            for c in row
        ])
    t = Table(data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(NAVY)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(LIGHT_BG)]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return t


def pdf_action_cards(story: list, actions: list[dict[str, str]], styles: dict) -> None:
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import KeepTogether, Paragraph, Spacer, Table, TableStyle

    if not actions:
        return

    usable = 17 * cm
    for act in actions[:8]:
        tier = act.get("priority_tier", "P3")
        tier_color = TIER_COLORS.get(tier, SLATE)
        rank = act.get("rank", "")
        issue = _esc(act.get("issue", ""))
        location = _esc(act.get("location", ""))
        platform = _esc(act.get("platform", ""))
        score = _esc(act.get("score", ""))
        action = _esc(act.get("action", ""))
        team = _esc(act.get("team", ""))
        due = _esc(act.get("due", ""))
        evidence = _esc(act.get("evidence", ""))

        header = Table(
            [[
                Paragraph(f"<font color='{tier_color}'><b>{tier}</b></font>", styles["cell_bold"]),
                Paragraph(f"<b>#{rank} {issue}</b>", styles["cell_bold"]),
            ]],
            colWidths=[1.2 * cm, usable - 1.2 * cm],
        )
        header.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(LIGHT_BG)),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))

        meta = Table(
            [[
                Paragraph(f"<b>Lokasi:</b> {location}", styles["cell"]),
                Paragraph(f"<b>Platform:</b> {platform}", styles["cell"]),
            ], [
                Paragraph(f"<b>Skor:</b> {score}", styles["cell"]),
                Paragraph(f"<b>Tarikh:</b> {due}", styles["cell"]),
            ], [
                Paragraph(f"<b>Pasukan:</b> {team}", styles["cell"]),
                Paragraph(f"<b>Bukti:</b> {evidence}", styles["cell"]),
            ]],
            colWidths=[usable / 2, usable / 2],
        )
        meta.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#e2e8f0")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))

        action_row = Table(
            [[Paragraph(f"<b>Tindakan:</b> {action}", styles["cell"])]],
            colWidths=[usable],
        )
        action_row.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.white),
            ("BOX", (0, 0), (-1, -1), 0.3, colors.HexColor("#e2e8f0")),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))

        card = Table([[header], [meta], [action_row]], colWidths=[usable])
        card.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor(NAVY)),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(KeepTogether([card, Spacer(1, 0.25 * cm)]))


def pdf_page_template(state: str, short: str):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm

    def _draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor(SLATE))
        canvas.drawString(2 * cm, A4[1] - 1.2 * cm, f"InsightPulse · PRN Naratif Cina · {state} ({short})")
        canvas.drawRightString(A4[0] - 2 * cm, 1 * cm, f"Muka {doc.page}")
        canvas.setStrokeColor(colors.HexColor("#e2e8f0"))
        canvas.setLineWidth(0.5)
        canvas.line(2 * cm, A4[1] - 1.4 * cm, A4[0] - 2 * cm, A4[1] - 1.4 * cm)
        canvas.restoreState()

    return _draw


# ── DOCX helpers ─────────────────────────────────────────────────────────────

def docx_set_cell_shading(cell, hex_color: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), hex_color.lstrip("#"))
    cell._tc.get_or_add_tcPr().append(shading)


def docx_style_table_header(row_cells, bg: str = "1e3a5f") -> None:
    from docx.shared import Pt, RGBColor

    for cell in row_cells:
        docx_set_cell_shading(cell, bg)
        for p in cell.paragraphs:
            if p.runs:
                for run in p.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.bold = True
                    run.font.size = Pt(9)
            else:
                txt = p.text
                p.text = ""
                run = p.add_run(txt)
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.bold = True
                run.font.size = Pt(9)


def docx_set_col_widths(table, widths_in: list[float]) -> None:
    from docx.shared import Inches

    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths_in):
                cell.width = Inches(widths_in[i])


def docx_action_cards(doc, actions: list[dict[str, str]]) -> None:
    from docx.shared import Pt, RGBColor

    if not actions:
        return

    for act in actions[:8]:
        tier = act.get("priority_tier", "P3")
        p = doc.add_paragraph()
        run = p.add_run(f"{tier}  ")
        run.bold = True
        run.font.size = Pt(11)
        tier_rgb = {"P1": RGBColor(220, 38, 38), "P2": RGBColor(245, 158, 11), "P3": RGBColor(37, 99, 235)}
        run.font.color.rgb = tier_rgb.get(tier, RGBColor(100, 116, 139))
        run2 = p.add_run(f"#{act.get('rank', '')}  {act.get('issue', '')}")
        run2.bold = True
        run2.font.size = Pt(11)

        t = doc.add_table(rows=3, cols=2)
        t.style = "Table Grid"
        fields = [
            ("Lokasi", act.get("location", "")),
            ("Platform", act.get("platform", "")),
            ("Skor", act.get("score", "")),
            ("Tarikh", act.get("due", "")),
            ("Pasukan", act.get("team", "")),
            ("Bukti", act.get("evidence", "")),
        ]
        for i, (label, val) in enumerate(fields):
            r, c = divmod(i, 2)
            cell = t.rows[r].cells[c]
            cell.text = f"{label}: {val}"
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)

        doc.add_paragraph(f"Tindakan: {act.get('action', '')}").runs[0].font.size = Pt(9)
        doc.add_paragraph("")


def docx_wrapped_table(doc, headers: list[str], rows: list[list[str]], widths_in: list[float]) -> None:
    from docx.shared import Pt

    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    docx_set_col_widths(t, widths_in)
    for j, h in enumerate(headers):
        t.rows[0].cells[j].text = h
    docx_style_table_header(t.rows[0].cells)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            cell.text = str(val)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)
