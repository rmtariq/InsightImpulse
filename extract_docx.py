#!/usr/bin/env python3
"""Extract and display Word document content"""

from docx import Document
import sys

def extract_docx(filepath):
    """Extract all text from a Word document"""
    try:
        doc = Document(filepath)
        
        print("=" * 80)
        print(f"📄 DOCUMENT: {filepath.split('/')[-1]}")
        print("=" * 80)
        print()
        
        # Extract all paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                print(para.text)
        
        # Extract tables if any
        if doc.tables:
            print("\n" + "=" * 80)
            print("📊 TABLES")
            print("=" * 80)
            for i, table in enumerate(doc.tables):
                print(f"\n--- Table {i+1} ---")
                for row in table.rows:
                    row_text = " | ".join([cell.text for cell in row.cells])
                    print(row_text)
        
        print("\n" + "=" * 80)
        print(f"✅ Total paragraphs: {len(doc.paragraphs)}")
        print(f"✅ Total tables: {len(doc.tables)}")
        print("=" * 80)
        
    except Exception as e:
        print(f"❌ Error: {e}")
        sys.exit(1)

if __name__ == "__main__":
    filepath = sys.argv[1] if len(sys.argv) > 1 else "/Users/rmtariq/Documents/InsightPulse/JITP_2026/Kalau tuan mahu, saya boleh sambung terus dan tuli.docx"
    extract_docx(filepath)
