#!/usr/bin/env python3
"""
Convert InsightPulse Markdown documentation to DOCX format
Generates professional Word documents with proper formatting
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import re
from pathlib import Path

def create_styled_document():
    """Create a document with custom styles"""
    doc = Document()
    
    # Set up styles
    styles = doc.styles
    
    # Heading 1 style
    h1 = styles['Heading 1']
    h1.font.name = 'Arial'
    h1.font.size = Pt(24)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 51, 102)  # Dark blue
    
    # Heading 2 style
    h2 = styles['Heading 2']
    h2.font.name = 'Arial'
    h2.font.size = Pt(18)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 102, 204)  # Medium blue
    
    # Heading 3 style
    h3 = styles['Heading 3']
    h3.font.name = 'Arial'
    h3.font.size = Pt(14)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(51, 102, 153)  # Light blue
    
    # Normal style
    normal = styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    
    return doc

def parse_markdown_to_docx(md_file, output_file):
    """Convert Markdown to DOCX with formatting"""
    
    print(f"📄 Converting {md_file} to DOCX...")
    
    # Read markdown file
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Create document
    doc = create_styled_document()
    
    # Split into lines
    lines = content.split('\n')
    
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        # Skip empty lines
        if not line:
            i += 1
            continue
        
        # Title (# at start)
        if line.startswith('# ') and not line.startswith('## '):
            title = line.replace('# ', '')
            p = doc.add_heading(title, level=1)
            
        # Heading 2 (##)
        elif line.startswith('## ') and not line.startswith('### '):
            heading = line.replace('## ', '')
            doc.add_heading(heading, level=2)
            
        # Heading 3 (###)
        elif line.startswith('### ') and not line.startswith('#### '):
            heading = line.replace('### ', '')
            doc.add_heading(heading, level=3)
            
        # Heading 4 (####)
        elif line.startswith('#### '):
            heading = line.replace('#### ', '')
            doc.add_heading(heading, level=4)
            
        # Horizontal rule
        elif line.startswith('---'):
            doc.add_paragraph('_' * 80)
            
        # Bullet list
        elif line.startswith('- ') or line.startswith('* '):
            text = line[2:].strip()
            # Remove markdown formatting
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)  # Bold
            text = re.sub(r'`(.*?)`', r'\1', text)  # Code
            p = doc.add_paragraph(text, style='List Bullet')
            
        # Numbered list
        elif re.match(r'^\d+\.\s', line):
            text = re.sub(r'^\d+\.\s', '', line)
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            text = re.sub(r'`(.*?)`', r'\1', text)
            p = doc.add_paragraph(text, style='List Number')
            
        # Code block (```)
        elif line.startswith('```'):
            # Skip the opening ```
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            
            # Add code block
            if code_lines:
                code_text = '\n'.join(code_lines)
                p = doc.add_paragraph(code_text)
                p.style = 'Normal'
                p_format = p.paragraph_format
                p_format.left_indent = Inches(0.5)
                # Set monospace font
                for run in p.runs:
                    run.font.name = 'Courier New'
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(64, 64, 64)
            
        # Table (|)
        elif line.startswith('|'):
            # Collect table rows
            table_rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = [cell.strip() for cell in lines[i].split('|')[1:-1]]
                # Skip separator rows (contains ---)
                if not any('---' in cell for cell in row):
                    table_rows.append(row)
                i += 1
            i -= 1  # Back up one
            
            if table_rows:
                # Create table
                table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
                table.style = 'Light Grid Accent 1'
                
                # Fill table
                for row_idx, row_data in enumerate(table_rows):
                    for col_idx, cell_data in enumerate(row_data):
                        cell = table.rows[row_idx].cells[col_idx]
                        # Remove markdown formatting
                        cell_data = re.sub(r'\*\*(.*?)\*\*', r'\1', cell_data)
                        cell_data = re.sub(r'`(.*?)`', r'\1', cell_data)
                        cell.text = cell_data
                        
                        # Bold header row
                        if row_idx == 0:
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    run.font.bold = True
        
        # Regular paragraph
        else:
            # Remove markdown formatting
            text = re.sub(r'\*\*(.*?)\*\*', r'\1', line)  # Bold
            text = re.sub(r'`(.*?)`', r'\1', text)  # Code
            text = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', text)  # Links
            
            # Skip if line is just formatting
            if text and text not in ['---', '===']:
                p = doc.add_paragraph(text)
        
        i += 1
    
    # Save document
    doc.save(output_file)
    print(f"✅ Saved: {output_file}")
    
    return output_file

def main():
    """Convert all documentation files"""
    
    files_to_convert = [
        ('INSIGHTPULSE_WHITE_PAPER.md', 'INSIGHTPULSE_WHITE_PAPER.docx'),
        ('INSIGHTPULSE_COMPLETE_GUIDE.md', 'INSIGHTPULSE_COMPLETE_GUIDE.docx'),
        ('QUICK_START_FOR_COLLEAGUES.md', 'QUICK_START_FOR_COLLEAGUES.docx'),
        ('DATA_STORAGE_STRUCTURE.md', 'DATA_STORAGE_STRUCTURE.docx'),
        ('CSV_STORAGE_SUMMARY.md', 'CSV_STORAGE_SUMMARY.docx'),
    ]
    
    print("="*80)
    print("📝 CONVERTING MARKDOWN TO DOCX")
    print("="*80)
    
    converted = []
    
    for md_file, docx_file in files_to_convert:
        if Path(md_file).exists():
            try:
                output = parse_markdown_to_docx(md_file, docx_file)
                converted.append(output)
            except Exception as e:
                print(f"❌ Error converting {md_file}: {e}")
        else:
            print(f"⚠️  File not found: {md_file}")
    
    print("\n" + "="*80)
    print(f"✅ CONVERSION COMPLETE!")
    print(f"   Converted {len(converted)} files")
    print("="*80)
    
    for file in converted:
        size = Path(file).stat().st_size / 1024
        print(f"   📄 {file} ({size:.1f} KB)")

if __name__ == "__main__":
    main()
