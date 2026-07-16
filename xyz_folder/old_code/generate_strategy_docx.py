#!/usr/bin/env python3
"""
Generate DOCX document for Crawling Strategy
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_heading_with_color(doc, text, level, color_rgb=(0, 102, 204)):
    """Add colored heading"""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(*color_rgb)
    return heading

def add_table_with_style(doc, data, headers):
    """Add styled table"""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Light Grid Accent 1'
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        # Bold header
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    # Data rows
    for row_data in data:
        row_cells = table.add_row().cells
        for i, cell_data in enumerate(row_data):
            row_cells[i].text = str(cell_data)
    
    return table

def create_strategy_document():
    """Create comprehensive strategy DOCX"""
    
    doc = Document()
    
    # Title
    title = doc.add_heading('InsightPulse Crawling Strategy', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Post + Comments Data Collection Strategy')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.color.rgb = RGBColor(128, 128, 128)
    
    doc.add_paragraph()  # Spacing
    
    # ========================================
    # 1. RATIO STRATEGY
    # ========================================
    add_heading_with_color(doc, '1. RATIO STRATEGY (30:70)', 1)
    
    doc.add_paragraph('Posts : Comments = 30% : 70%')
    
    # Table: Dataset sizes
    headers = ['Dataset Size', 'Posts', 'Comments', 'Comments/Post', 'Cost (X/Twitter)']
    data = [
        ['Quick Test (50)', '15', '35', '2-3', '$0.09'],
        ['Quick Analysis (1,000)', '300', '700', '2-3', '$1.83'],
        ['Standard Analysis (5,000)', '1,500', '3,500', '2-3', '$9.13'],
        ['Deep Analysis (10,000)', '3,000', '7,000', '2-3', '$18.25'],
        ['Market Research (25,000)', '7,500', '17,500', '2-3', '$45.63'],
        ['Enterprise Analysis (50,000)', '15,000', '35,000', '2-3', '$91.25'],
        ['Big Data Insights (100,000)', '30,000', '70,000', '2-3', '$182.50'],
    ]
    add_table_with_style(doc, data, headers)
    
    doc.add_paragraph()
    
    # Rationale
    p = doc.add_paragraph()
    p.add_run('Rationale:').bold = True
    doc.add_paragraph('✅ Posts = Data utama (topik, context)', style='List Bullet')
    doc.add_paragraph('✅ Comments = Data sokongan (insights, sentiment, opinions)', style='List Bullet')
    doc.add_paragraph('✅ Comments lebih penting untuk analytics mendalam', style='List Bullet')
    doc.add_paragraph('✅ Public opinion terdapat dalam comments', style='List Bullet')
    
    doc.add_page_break()
    
    # ========================================
    # 2. CRAWLING WORKFLOW
    # ========================================
    add_heading_with_color(doc, '2. CRAWLING WORKFLOW', 1)
    
    doc.add_paragraph('Sequential Approach: Posts → Comments')
    
    workflow = doc.add_paragraph()
    workflow.add_run('STEP 1: Crawl Posts\n').bold = True
    workflow.add_run('   • Use AI-generated keywords\n')
    workflow.add_run('   • Target: 30% of dataset size\n')
    workflow.add_run('   • Filter by engagement, date range\n\n')
    
    workflow.add_run('STEP 2: Analyze & Rank Posts\n').bold = True
    workflow.add_run('   • Rank by engagement (likes + shares + comments)\n')
    workflow.add_run('   • Filter spam, duplicates\n')
    workflow.add_run('   • Select quality posts\n\n')
    
    workflow.add_run('STEP 3: Crawl Comments\n').bold = True
    workflow.add_run('   • Use post URLs from Step 1\n')
    workflow.add_run('   • Target: 70% of dataset size\n')
    workflow.add_run('   • Equal distribution per post\n\n')
    
    workflow.add_run('STEP 4: Combine & Process\n').bold = True
    workflow.add_run('   • Merge posts + comments\n')
    workflow.add_run('   • Sentiment analysis\n')
    workflow.add_run('   • Generate reports\n')
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.add_run('Advantages:').bold = True
    doc.add_paragraph('✅ Faster execution (parallel post crawling)', style='List Bullet')
    doc.add_paragraph('✅ Cost efficient (filter before comments)', style='List Bullet')
    doc.add_paragraph('✅ Better data quality', style='List Bullet')
    doc.add_paragraph('✅ Easier error handling', style='List Bullet')
    doc.add_paragraph('✅ Flexible adjustment', style='List Bullet')
    
    doc.add_page_break()
    
    # ========================================
    # 3. SAMPLING STRATEGY
    # ========================================
    add_heading_with_color(doc, '3. SAMPLING STRATEGY', 1)
    
    doc.add_paragraph('Equal Sampling (All Posts Get Same Comments)')
    
    p = doc.add_paragraph()
    p.add_run('Formula:\n').bold = True
    p.add_run('comments_per_post = total_comments_target ÷ total_posts')
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.add_run('Example (1,000 results):\n').bold = True
    p.add_run('Posts: 300\n')
    p.add_run('Comments target: 700\n')
    p.add_run('Comments per post: 700 ÷ 300 = 2-3 comments per post')
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.add_run('Implementation:').bold = True
    doc.add_paragraph('• All posts treated equally', style='List Bullet')
    doc.add_paragraph('• Simple and fair distribution', style='List Bullet')
    doc.add_paragraph('• Predictable results', style='List Bullet')
    doc.add_paragraph('• Easy to calculate costs', style='List Bullet')
    
    doc.add_page_break()
    
    # ========================================
    # 4. ADAPTIVE ALGORITHM
    # ========================================
    add_heading_with_color(doc, '4. ADAPTIVE ALGORITHM', 1)
    
    doc.add_paragraph('Scenario-Based Adjustment')
    
    # Scenario 1
    add_heading_with_color(doc, 'Scenario 1: Banyak Posts, Sikit Comments', 2, (255, 102, 0))

    p = doc.add_paragraph()
    p.add_run('Condition: ').bold = True
    p.add_run('avg_comments_available < 3')

    p = doc.add_paragraph()
    p.add_run('Strategy: MAXIMIZE POSTS\n').bold = True
    p.add_run('• Increase posts allocation up to 50%\n')
    p.add_run('• Take available comments (1-2 per post)\n')
    p.add_run('• Prioritize post diversity\n')

    p = doc.add_paragraph()
    p.add_run('Example:\n').bold = True
    p.add_run('Target: 1,000 results (300 posts + 700 comments)\n')
    p.add_run('Found: 500 posts, avg 1.5 comments/post\n')
    p.add_run('Action: Take 400 posts + 600 comments = 1,000 ✅')

    doc.add_paragraph()

    # Scenario 2
    add_heading_with_color(doc, 'Scenario 2: Sikit Posts, Banyak Comments', 2, (255, 102, 0))

    p = doc.add_paragraph()
    p.add_run('Condition: ').bold = True
    p.add_run('posts_found < target_posts')

    p = doc.add_paragraph()
    p.add_run('Strategy: MAXIMIZE COMMENTS\n').bold = True
    p.add_run('• Take all available posts\n')
    p.add_run('• Increase comments per post\n')
    p.add_run('• Cap at 50 comments per post\n')

    p = doc.add_paragraph()
    p.add_run('Example:\n').bold = True
    p.add_run('Target: 1,000 results (300 posts + 700 comments)\n')
    p.add_run('Found: 100 posts, avg 50 comments/post\n')
    p.add_run('Action: Take 100 posts + 900 comments = 1,000 ✅')

    doc.add_paragraph()

    # Scenario 3
    add_heading_with_color(doc, 'Scenario 3: Normal (Ideal Case)', 2, (0, 153, 0))

    p = doc.add_paragraph()
    p.add_run('Condition: ').bold = True
    p.add_run('posts_found >= target_posts AND avg_comments >= 3')

    p = doc.add_paragraph()
    p.add_run('Strategy: FOLLOW 30:70 RATIO\n').bold = True
    p.add_run('• Take target posts (30%)\n')
    p.add_run('• Take target comments (70%)\n')
    p.add_run('• Standard distribution\n')

    p = doc.add_paragraph()
    p.add_run('Example:\n').bold = True
    p.add_run('Target: 1,000 results\n')
    p.add_run('Found: 500 posts, avg 5 comments/post\n')
    p.add_run('Action: Take 300 posts + 700 comments = 1,000 ✅')

    doc.add_page_break()

    # ========================================
    # 5. SMART FILTERING
    # ========================================
    add_heading_with_color(doc, '5. SMART FILTERING RULES', 1)

    add_heading_with_color(doc, 'Post Selection Criteria (Priority Order):', 2, (0, 102, 204))

    p = doc.add_paragraph()
    p.add_run('1. High Engagement\n').bold = True
    p.add_run('   • Total engagement = likes + shares + comments + views\n')
    p.add_run('   • Rank posts by total engagement\n')
    p.add_run('   • Select top N posts\n\n')

    p.add_run('2. Date Relevance\n').bold = True
    p.add_run('   • Within specified date range\n')
    p.add_run('   • Prefer recent posts for trending topics\n')
    p.add_run('   • Historical posts for trend analysis\n\n')

    p.add_run('3. Content Quality\n').bold = True
    p.add_run('   • Minimum 20 words\n')
    p.add_run('   • Not spam or promotional\n')
    p.add_run('   • Has meaningful content\n')
    p.add_run('   • Diverse authors\n\n')

    p.add_run('4. Diversity\n').bold = True
    p.add_run('   • Avoid multiple posts from same author\n')
    p.add_run('   • Different perspectives\n')
    p.add_run('   • Various post types (original, retweet, quote)')

    doc.add_paragraph()

    add_heading_with_color(doc, 'Comment Selection Criteria (Priority Order):', 2, (0, 102, 204))

    p = doc.add_paragraph()
    p.add_run('1. Top Comments\n').bold = True
    p.add_run('   • Most liked/upvoted\n')
    p.add_run('   • High engagement\n')
    p.add_run('   • Quality responses\n\n')

    p.add_run('2. Recent Comments\n').bold = True
    p.add_run('   • Latest discussions\n')
    p.add_run('   • Current sentiment\n')
    p.add_run('   • Trending opinions\n\n')

    p.add_run('3. Content Quality\n').bold = True
    p.add_run('   • Minimum 10 words\n')
    p.add_run('   • Substantial content\n')
    p.add_run('   • Not spam or bot-generated\n\n')

    p.add_run('4. Diversity\n').bold = True
    p.add_run('   • Different authors\n')
    p.add_run('   • Various viewpoints\n')
    p.add_run('   • Balanced sentiment')

    doc.add_page_break()

    # ========================================
    # 6. FALLBACK RULES
    # ========================================
    add_heading_with_color(doc, '6. FALLBACK RULES', 1)

    p = doc.add_paragraph()
    p.add_run('Rule 1: Minimum Posts Guarantee\n').bold = True
    p.add_run('minimum_posts = dataset_size × 0.10\n')
    p.add_run('Always ensure at least 10% posts, even if comments are abundant\n\n')

    p.add_run('Rule 2: Maximum Comments per Post\n').bold = True
    p.add_run('max_comments_per_post = 50\n')
    p.add_run('If post has >50 comments:\n')
    p.add_run('   • Take top 25 (most liked)\n')
    p.add_run('   • Take recent 25 (latest)\n\n')

    p.add_run('Rule 3: Quality Threshold\n').bold = True
    p.add_run('Skip if:\n')
    p.add_run('   • Text length < 5 words\n')
    p.add_run('   • Spam detected\n')
    p.add_run('   • Duplicate content\n')
    p.add_run('   • Bot-generated\n')
    p.add_run('   • Invalid/deleted content\n\n')

    p.add_run('Rule 4: Minimum Data Guarantee\n').bold = True
    p.add_run('If total_results < target × 0.50:\n')
    p.add_run('   • Expand search keywords\n')
    p.add_run('   • Extend date range\n')
    p.add_run('   • Try related hashtags\n')
    p.add_run('   • Use alternative actors')

    doc.add_page_break()

    # ========================================
    # 7. MULTI-PLATFORM
    # ========================================
    add_heading_with_color(doc, '7. MULTI-PLATFORM DISTRIBUTION', 1)

    doc.add_paragraph('Equal Split Strategy')

    p = doc.add_paragraph()
    p.add_run('Formula:\n').bold = True
    p.add_run('per_platform_allocation = total_dataset_size ÷ number_of_platforms')

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.add_run('Example: 1,000 results, 3 platforms (X, Facebook, Instagram)\n\n').bold = True
    p.add_run('Per platform: 1,000 ÷ 3 = 333 results\n\n')
    p.add_run('X (Twitter):\n')
    p.add_run('   • Posts: 100 (30%)\n')
    p.add_run('   • Comments: 233 (70%)\n')
    p.add_run('   • Total: 333\n\n')
    p.add_run('Facebook:\n')
    p.add_run('   • Posts: 100 (30%)\n')
    p.add_run('   • Comments: 233 (70%)\n')
    p.add_run('   • Total: 333\n\n')
    p.add_run('Instagram:\n')
    p.add_run('   • Posts: 100 (30%)\n')
    p.add_run('   • Comments: 233 (70%)\n')
    p.add_run('   • Total: 333\n\n')
    p.add_run('Grand Total: 300 posts + 699 comments ≈ 1,000 ✅')

    doc.add_page_break()

    # ========================================
    # 8. PLATFORM-SPECIFIC
    # ========================================
    add_heading_with_color(doc, '8. PLATFORM-SPECIFIC STRATEGIES', 1)

    # X/Twitter
    add_heading_with_color(doc, 'X (Twitter)', 2, (29, 161, 242))
    p = doc.add_paragraph()
    p.add_run('Posts Actor: ').bold = True
    p.add_run('kaitoeasyapi/twitter-x-data-tweet-scraper-pay-per-result-cheapest\n')
    p.add_run('Comments Actor: ').bold = True
    p.add_run('scraper_one/x-post-replies-scraper\n\n')
    p.add_run('Strategy:\n').bold = True
    p.add_run('• Use searchTerms parameter (array)\n')
    p.add_run('• Include retweets and quotes\n')
    p.add_run('• Filter by engagement\n')
    p.add_run('• Get replies using tweet URLs')

    doc.add_paragraph()

    # Instagram
    add_heading_with_color(doc, 'Instagram', 2, (225, 48, 108))
    p = doc.add_paragraph()
    p.add_run('Posts Actor: ').bold = True
    p.add_run('apify/instagram-scraper\n')
    p.add_run('Comments Actor: ').bold = True
    p.add_run('Same actor with post URLs\n\n')
    p.add_run('Strategy:\n').bold = True
    p.add_run('• Search by hashtags\n')
    p.add_run('• Filter by engagement rate\n')
    p.add_run('• Get comments from post URLs\n')
    p.add_run('• Include story mentions')

    doc.add_paragraph()

    # Facebook
    add_heading_with_color(doc, 'Facebook', 2, (24, 119, 242))
    p = doc.add_paragraph()
    p.add_run('Posts Actor: ').bold = True
    p.add_run('apify/facebook-posts-scraper\n')
    p.add_run('Comments Actor: ').bold = True
    p.add_run('Same actor with post URLs\n\n')
    p.add_run('Strategy:\n').bold = True
    p.add_run('• Search by keywords in groups/pages\n')
    p.add_run('• Filter by reactions + shares\n')
    p.add_run('• Get comments and replies\n')
    p.add_run('• Include reactions data')

    doc.add_paragraph()

    # Save document
    filename = 'InsightPulse_Crawling_Strategy.docx'
    doc.save(filename)
    print(f"✅ Document created: {filename}")
    return filename

if __name__ == "__main__":
    create_strategy_document()

