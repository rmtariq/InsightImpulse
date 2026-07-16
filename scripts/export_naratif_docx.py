#!/usr/bin/env python3
"""Eksport DOCX cadangan tindakan + kit copywriting — Johor & N9 berasingan."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Dict, List, Optional

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "data/projects/political/PRN/reports/naratif_cina_india"
PLAN_JSON = OUT_DIR / "naratif_cadangan_tindakan.json"
KIT_JSON = OUT_DIR / "naratif_action_kit.json"


def _rgb(doc_module, r: int, g: int, b: int):
    from docx.shared import RGBColor
    return RGBColor(r, g, b)


def _heading(doc, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _bullet_list(doc, items: List[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _action_block(doc, act: Dict[str, Any]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc.add_page_break()
    h = doc.add_heading(f"{act['priority']}: {act['title']}", level=1)

    meta = doc.add_paragraph()
    meta.add_run(f"ID: {act['id']}  |  Komuniti: {act['community'].upper()}  |  Isu: {act['issue']}")
    meta.add_run(f"\nDUN sasaran: {', '.join(act['duns'])}")
    meta.add_run(f"\nPemilik: {act['owner']}  |  Tarikh akhir: {act['deadline']}  |  Saluran: {act['channel']}")

    _heading(doc, "1. Apa yang perlu dilakukan", 2)
    doc.add_paragraph(act["what"])

    _heading(doc, "2. Mengapa ini penting (berdasarkan intel crawl)", 2)
    doc.add_paragraph(act["why"])

    _heading(doc, "3. Bagaimana melaksanakan — langkah demi langkah", 2)
    _bullet_list(doc, act.get("how_steps", []))

    _heading(doc, "4. Mesej yang hendak disampaikan (salin & edit)", 2)
    msg = act.get("message", {})
    if msg.get("bm"):
        p = doc.add_paragraph()
        p.add_run("BM: ").bold = True
        p.add_run(msg["bm"])
    if msg.get("zh"):
        p = doc.add_paragraph()
        p.add_run("中文: ").bold = True
        p.add_run(msg["zh"])
    if msg.get("ta"):
        p = doc.add_paragraph()
        p.add_run("தமிழ்: ").bold = True
        p.add_run(msg["ta"])

    _heading(doc, "5. KPI / ukuran kejayaan", 2)
    doc.add_paragraph(act.get("kpi", "—"))


def _dun_table(doc, rows: List[Dict[str, str]], state_label: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc.add_page_break()
    _heading(doc, f"Jadual Cadangan Semua DUN — {state_label}", 1)
    doc.add_paragraph(
        "Setiap kerusi DUN di bawah mempunyai cadangan tindakan minimum untuk komuniti Cina dan India. "
        "Gunakan sebagai checklist operasi calon & ground team. P1 = keutamaan tinggi berdasarkan profil demografi."
    )

    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ["DUN", "Nama", "Keutamaan", "Isu utama", "Tindakan Cina", "Tindakan India"]
    for i, hdr in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = hdr
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True

    for row in rows:
        cells = table.add_row().cells
        cells[0].text = row["dun_code"]
        cells[1].text = row["dun_name"]
        cells[2].text = row["priority"]
        cells[3].text = row["primary_issue"]
        cells[4].text = row["chinese_action"]
        cells[5].text = row["indian_action"]


def _copy_kit_section(doc, kit: Dict[str, Any], state_key: str, label: str) -> None:
    st = kit.get(state_key, {})
    if not st:
        return
    doc.add_page_break()
    _heading(doc, f"Kit Copywriting Socmed — {label}", 1)

    for comm, title in [("chinese", "Komuniti Cina"), ("indian", "Komuniti India"), ("punjabi", "Komuniti Punjabi")]:
        c = st.get(comm, {})
        if not c:
            continue
        _heading(doc, title, 2)
        if c.get("note"):
            doc.add_paragraph(c["note"])

        for i, post in enumerate(c.get("facebook", []), 1):
            doc.add_paragraph(f"Post #{i}: {post.get('title', '')}", style="List Number")
            if post.get("bm"):
                doc.add_paragraph(f"BM: {post['bm']}")
            if post.get("zh"):
                doc.add_paragraph(f"中文: {post['zh']}")
            if post.get("ta"):
                doc.add_paragraph(f"தமிழ்: {post['ta']}")
            if post.get("en"):
                doc.add_paragraph(f"EN: {post['en']}")
            if post.get("hashtags"):
                doc.add_paragraph(f"Hashtag: {post['hashtags']}")
            if post.get("cta"):
                doc.add_paragraph(f"CTA: {post['cta']}")

        poster = c.get("poster", {})
        if poster:
            _heading(doc, "Teks Poster", 3)
            doc.add_paragraph(f"Headline BM: {poster.get('headline_bm', '')}")
            doc.add_paragraph(f"Headline 中文/Tamil: {poster.get('headline_zh', poster.get('headline_ta', ''))}")
            doc.add_paragraph(f"Sub BM: {poster.get('sub_bm', '')}")
            doc.add_paragraph(f"Footer: {poster.get('footer', '')}")

        tiktok = c.get("tiktok_prompt") or c.get("tiktok_script_30s")
        if tiktok:
            _heading(doc, "Prompt Video Pendek (TikTok/CapCut)", 3)
            doc.add_paragraph(tiktok)


def _counter_section(doc, kit: Dict[str, Any]) -> None:
    counters = kit.get("counter_narratives", {})
    if not counters:
        return
    doc.add_page_break()
    _heading(doc, "Counter-Narrative (Copy-paste)", 1)
    for key, v in counters.items():
        if not isinstance(v, dict):
            continue
        doc.add_paragraph(key.replace("_", " ").title(), style="List Bullet")
        for lang, txt in v.items():
            if lang != "do_not":
                doc.add_paragraph(f"  {lang}: {txt}")
        if v.get("do_not"):
            p = doc.add_paragraph()
            p.add_run(f"⚠ JANGAN: {v['do_not']}").bold = True


def write_state_docx(
    state_key: str,
    state_label: str,
    plan: Dict[str, Any],
    kit: Optional[Dict[str, Any]] = None,
    out_path: Optional[Path] = None,
) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    navy = _rgb(None, 30, 58, 95)

    # Cover
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("InsightPulse · PAS / PN / MN")
    r.bold = True
    r.font.size = Pt(14)

    h = doc.add_heading(f"Cadangan Tindakan Naratif Cina & India\n{state_label}", level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for line in [
        f"Negeri: {state_label}",
        f"Dijana: {plan.get('generated_at', '')}",
        "Klasifikasi: INTERNAL — Kempen PRN",
        plan.get("coalition", ""),
    ]:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        "Dokumen ini mengandungi cadangan tindakan operasi yang terperinci — "
        "bukan ringkasan. Setiap tindakan menerangkan apa, mengapa, bagaimana, dan mesej yang hendak disampaikan.",
        style="Intense Quote",
    )

    # Executive
    doc.add_page_break()
    _heading(doc, "Ringkasan Eksekutif", 1)
    if plan.get("executive_summary"):
        doc.add_paragraph(plan["executive_summary"])
    st_data = plan.get(state_key, {})
    if st_data.get("stats_note"):
        doc.add_paragraph(st_data["stats_note"])

    _heading(doc, "Peraturan Nada (WAJIB)", 2)
    _bullet_list(doc, plan.get("tone_rules", []))

    _heading(doc, "Keutamaan Koalisi", 2)
    _bullet_list(doc, plan.get("coalition_priorities", []) or ["—"])

    _heading(doc, "JANGAN Buat", 2)
    _bullet_list(doc, plan.get("do_not_do", []))

    # Priority actions
    doc.add_page_break()
    _heading(doc, f"Tindakan Keutamaan — {state_label}", 1)
    doc.add_paragraph(
        "Bahagian ini mengandungi tindakan P1 (mesti minggu ini) dan P2 (susulan 7–14 hari). "
        "Setiap kad mempunyai langkah operasi yang boleh diagihkan kepada calon, exco, dan ground team."
    )
    for act in st_data.get("priority_actions", []):
        _action_block(doc, act)

    # DUN matrix
    _dun_table(doc, st_data.get("dun_matrix", []), state_label)

    # Copy kit
    if kit:
        kit_key = "johor" if state_key == "johor" else "n9"
        _copy_kit_section(doc, kit, kit_key, state_label)
        _counter_section(doc, kit)

        videos = kit.get("video_prompts_capcut", [])
        if videos:
            doc.add_page_break()
            _heading(doc, "Prompt Video Pendek (CapCut / TikTok)", 1)
            for v in videos:
                doc.add_paragraph(v.get("name", ""), style="List Bullet")
                doc.add_paragraph(v.get("prompt", ""))

    out = out_path or OUT_DIR / f"naratif_cadangan_tindakan_{state_key}.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def export_both(
    plan: Optional[Dict[str, Any]] = None,
    kit: Optional[Dict[str, Any]] = None,
) -> Dict[str, str]:
    if plan is None and PLAN_JSON.exists():
        plan = json.loads(PLAN_JSON.read_text(encoding="utf-8"))
    if kit is None and KIT_JSON.exists():
        kit = json.loads(KIT_JSON.read_text(encoding="utf-8"))
    if not plan:
        raise FileNotFoundError("Tiada pelan tindakan — jalankan analyze_naratif_agentic.py dahulu")

    paths = {}
    paths["johor"] = str(write_state_docx("johor", "Johor · 56 DUN", plan, kit))
    paths["n9"] = str(write_state_docx("n9", "Negeri Sembilan · 36 DUN", plan, kit))
    return paths


def main() -> int:
    paths = export_both()
    for k, p in paths.items():
        print(f"✅ DOCX {k}: {p}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
