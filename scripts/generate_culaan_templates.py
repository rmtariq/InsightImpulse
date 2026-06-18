#!/usr/bin/env python3
"""Generate DOCX + CSV templates for PRN Digital Culaan (N9, Johor, Melaka)."""
from __future__ import annotations

import csv
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from digital_cula_lib import (  # noqa: E402
    STATES,
    TEMPLATE_HEADERS,
    load_dun_catalog,
    template_csv_text,
)

OUT_DIR = ROOT / "data/projects/political/PRN/digital_cula"


def write_csv_templates() -> list[Path]:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []

    # Master template with examples
    master = OUT_DIR / "TEMPLATE_culaan.csv"
    master.write_text(template_csv_text(), encoding="utf-8-sig")
    written.append(master)

    # Empty template (headers only — for daily fill)
    empty = OUT_DIR / "TEMPLATE_culaan_KOSONG.csv"
    with empty.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(TEMPLATE_HEADERS)
    written.append(empty)

    examples = {
        "N9": [
            ["2026-06-17", "N25", "Paroi A", "PDM-Paroi", "Ketua PDM Paroi", "door_to_door", "20", "12", "3", "5", "2", "2", ""],
            ["2026-06-17", "N05", "Serting B", "PDM-Serting", "", "door_to_door", "15", "8", "2", "3", "2", "1", ""],
        ],
        "Johor": [
            ["2026-06-17", "N14", "Muar Pusat", "PDM-Muar", "", "door_to_door", "18", "10", "2", "4", "2", "2", ""],
            ["2026-06-17", "N48", "JB Tengah", "PDM-JB", "", "program", "25", "14", "4", "6", "2", "2", "Program malam"],
        ],
        "Melaka": [
            ["2026-06-17", "N03", "Klebang", "PDM-Klebang", "", "door_to_door", "12", "7", "1", "3", "2", "1", ""],
            ["2026-06-17", "N10", "Bachang", "PDM-Bachang", "", "follow_up", "8", "5", "1", "2", "1", "1", ""],
        ],
    }

    for state in STATES:
        path = OUT_DIR / f"TEMPLATE_culaan_{state}.csv"
        with path.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(TEMPLATE_HEADERS)
            for row in examples.get(state, []):
                w.writerow(row)
        written.append(path)

    return written


def write_docx_template() -> Path:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm, Pt, RGBColor
    except ImportError as e:
        raise SystemExit("python-docx required: pip install python-docx") from e

    doc = Document()
    navy = RGBColor(30, 58, 95)
    grey = RGBColor(80, 80, 80)

    # --- Cover ---
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("InsightPulse · PRN War Room")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = navy

    h = doc.add_heading("BORANG LAPORAN CULAAN DIGITAL", level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Negeri Sembilan · Johor · Melaka")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ver = doc.add_paragraph(f"Versi template: {datetime.now().strftime('%d %b %Y')}")
    ver.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        "NOTA: Tiada IC / alamat penuh pengundi. Borang ini untuk operasi jentera sahaja — bukan polling atau ramalan undi.",
        style="Intense Quote",
    )

    doc.add_heading("Cara Guna (Paling Cepat)", level=1)
    for line in [
        "Petugas isi jadual di bawah (Excel/Word) ATAU hantar format WhatsApp ke HQ.",
        "HQ semak offline → Save As CSV → Muat naik di War Room tab ⚡ Muat Naik Batch.",
        "Selepas upload, dashboard terus kemaskini (auto-disahkan).",
        "Fail CSV rakan: TEMPLATE_culaan_KOSONG.csv atau TEMPLATE_culaan_{N9|Johor|Melaka}.csv",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("Medan Wajib", level=1)
    fields = [
        ("tarikh_lawatan", "Tarikh lawatan", "2026-06-17 atau 17/06/2026"),
        ("kod_dun", "Kod DUN", "N25, N05, …"),
        ("pdm", "PDM / Kawasan", "Paroi A, Serting B, …"),
        ("pelapor_id", "ID pelapor", "PDM-Paroi / SV-001"),
        ("rumah_dilawati", "Bil. rumah dilawati", "Nombor"),
        ("cula_baharu", "Cula baharu hari ini", "Nombor (delta, bukan jumlah keseluruhan)"),
        ("cula_bulan / condong / pagar / dacing", "Pecahan status", "Optional — jumlah = cula baharu"),
        ("jenis_aktiviti", "Jenis", "door_to_door | program | follow_up | lain"),
        ("catatan", "Catatan ringkas", "Tanpa IC/alamat penuh"),
    ]
    ft = doc.add_table(rows=1, cols=3)
    ft.style = "Table Grid"
    hdr = ft.rows[0].cells
    hdr[0].text = "Medan CSV"
    hdr[1].text = "Label"
    hdr[2].text = "Contoh"
    for c in hdr:
        for p in c.paragraphs:
            for run in p.runs:
                run.bold = True
    for key, label, ex in fields:
        row = ft.add_row().cells
        row[0].text = key
        row[1].text = label
        row[2].text = ex

    doc.add_page_break()

    # --- Single entry form ---
    doc.add_heading("Borang Harian — Satu PDM (Hardcopy / Isi Manual)", level=1)
    doc.add_paragraph("Negeri: _______________   Tarikh: _______________   ID Pelapor: _______________")

    single = doc.add_table(rows=8, cols=2)
    single.style = "Table Grid"
    single_labels = [
        ("Kod DUN", ""),
        ("Nama DUN", ""),
        ("PDM / Kawasan", ""),
        ("Jenis aktiviti", "☐ door-to-door  ☐ program  ☐ follow-up  ☐ lain"),
        ("Rumah dilawati", ""),
        ("Cula baharu hari ini", ""),
        ("Pecahan: Bulan / Condong / Atas pagar / Dacing", "___ / ___ / ___ / ___"),
        ("Catatan", ""),
    ]
    for i, (label, val) in enumerate(single_labels):
        single.rows[i].cells[0].text = label
        single.rows[i].cells[1].text = val
        single.rows[i].cells[0].width = Cm(5)

    doc.add_paragraph("")
    p = doc.add_paragraph("Tandatangan Ketua PDM: _______________________   Tarikh: __________")
    p.runs[0].font.color.rgb = grey

    doc.add_page_break()

    # --- Batch table (10 rows) ---
    doc.add_heading("Borang Batch — Banyak PDM (Salin ke CSV)", level=1)
    doc.add_paragraph(
        "Isi baris di bawah. Di Excel: salin jadual → paste ke TEMPLATE_culaan_KOSONG.csv → upload."
    )

    batch_headers = [
        "Tarikh", "DUN", "PDM", "ID Pelapor", "Rumah", "Cula", "Bulan", "Condong", "Pagar", "Dacing", "Catatan"
    ]
    batch = doc.add_table(rows=11, cols=len(batch_headers))
    batch.style = "Table Grid"
    for j, hname in enumerate(batch_headers):
        batch.rows[0].cells[j].text = hname
        for run in batch.rows[0].cells[j].paragraphs[0].runs:
            run.bold = True
    for i in range(1, 11):
        batch.rows[i].cells[0].text = "____-__-__"

    doc.add_page_break()

    # --- WhatsApp format ---
    doc.add_heading("Format WhatsApp (Copy-Paste ke HQ)", level=1)
    doc.add_paragraph("Satu baris = satu laporan PDM. HQ tampal terus di tab Muat Naik Batch.")
    doc.add_paragraph(
        "Format:  DUN | PDM | rumah | cula | bulan | condong | pagar | dacing",
        style="Intense Quote",
    )
    doc.add_paragraph("Contoh:")
    for ex in [
        "N25 | Paroi A | 20 | 12 | 3 | 5 | 2 | 2",
        "N05 | Serting B | 15 | 8 | 2 | 3 | 2 | 1",
        "N14 | Muar Pusat | 18 | 10 | 2 | 4 | 2 | 2",
    ]:
        doc.add_paragraph(ex, style="List Bullet")

    doc.add_heading("Tukar Excel/Word → CSV", level=1)
    for step in [
        "Buka TEMPLATE_culaan_KOSONG.csv di Excel (atau copy jadual batch di atas).",
        "Isi satu baris per PDM.",
        "File → Save As → CSV UTF-8 (Comma delimited) *.csv",
        "War Room → Culaan Digital → ⚡ Muat Naik Batch → pilih fail → Upload.",
        "Semak tab 📊 Ringkasan DUN — nombor terus kemaskini.",
    ]:
        doc.add_paragraph(step, style="List Number")

    doc.add_heading("Rujukan Kod DUN (contoh)", level=1)
    for state in STATES:
        catalog = load_dun_catalog(state)
        doc.add_paragraph(f"{state} ({len(catalog)} kerusi):", style="List Bullet")
        sample = ", ".join(f"{k} {v}" for k, v in list(catalog.items())[:8])
        doc.add_paragraph(f"  {sample} …")

    out = OUT_DIR / "BORANG_CULAAN_DIGITAL_PRN.docx"
    doc.save(out)
    return out


def main() -> None:
    csv_files = write_csv_templates()
    docx = write_docx_template()
    print("Generated templates:")
    for p in csv_files:
        print(f"  CSV  {p.relative_to(ROOT)}")
    print(f"  DOCX {docx.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
